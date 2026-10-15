"""Dossier de relecture native NL / RO / PL (tableur Excel).

Lit src/locales/<langue>/*.json et produit docs/relecture/relecture-nl-ro-pl.xlsx :
  1. mentions légales des PDF et page de signature (priorité 1 et 2) ;
  2. écrans principaux : connexion, tableau de bord, sous-traitants, chantiers, contrats,
     prestations, conformité, exports.

Aucune donnée de base : uniquement les fichiers de traduction du dépôt.
Lancer (Python 3 + openpyxl) :  python scripts/relecture/exporter_relecture.py
La colonne « Clé technique » permet de réintégrer automatiquement les corrections.
"""

import json
import os
import re
from collections import OrderedDict
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

RACINE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
LOCALES = os.path.join(RACINE, "src", "locales")
SORTIE = os.path.join(RACINE, "docs", "relecture", "relecture-nl-ro-pl.xlsx")
LANGUES = ["nl", "ro", "pl"]
NOMS_LANGUES = {"nl": "Néerlandais (NL)", "ro": "Roumain (RO)", "pl": "Polonais (PL)"}
STATUTS = ["À valider", "Correction proposée", "Validé"]
SUFFIXES = ("_zero", "_one", "_two", "_few", "_many", "_other")


def charger(langue, ns):
    with open(os.path.join(LOCALES, langue, f"{ns}.json"), encoding="utf8") as f:
        return json.load(f, object_pairs_hook=OrderedDict)


def aplatir(d, prefixe=""):
    for k, v in d.items():
        cle = f"{prefixe}.{k}" if prefixe else k
        if isinstance(v, dict):
            yield from aplatir(v, cle)
        else:
            yield cle, v


def base_pluriel(cle):
    for s in SUFFIXES:
        if cle.endswith(s):
            return cle[: -len(s)], s[1:]
    return cle, None


def source_fr(fr, cle):
    """Texte français d'une clé ; pour une forme plurielle propre à RO/PL (few, many), la forme
    française « other » sert de référence."""
    if cle in fr:
        return fr[cle]
    base, forme = base_pluriel(cle)
    if forme:
        return fr.get(f"{base}_other") or fr.get(f"{base}_one") or ""
    return ""


# --- Priorités : mentions légales et textes destinés au client / sous-traitant ---------------
NOTE_PDF_RO_PL = (
    "Les documents clients (devis, contrats) ne sont produits qu'en FR, NL ou EN : en RO / PL, "
    "ce texte n'apparaît pas aujourd'hui sur un PDF (relecture moins urgente)."
)
PRIORITE_1 = OrderedDict(
    [
        ("pdf:mentionReverseCharge", ("PDF devis / facture — mentions légales",
          "Mention légale d'AUTOLIQUIDATION imprimée sur le devis quand le client est assujetti "
          "(régime du cocontractant, art. 20 AR n° 1 TVA). Doit reprendre la formulation légale "
          "attendue par l'administration dans la langue du client.")),
        ("pdf:mention6", ("PDF devis / facture — mentions légales",
          "Mention TVA 6 % (logement de plus de 10 ans, attestation du client requise). "
          "Formulation légale à vérifier.")),
        ("pdf:legalMentions", ("PDF devis / facture — mentions légales", "Titre de la rubrique des mentions légales.")),
        ("pdf:generalConditions", ("PDF devis — conditions", "Titre de la rubrique « Conditions » du devis.")),
        ("pdf:defaultConditions", ("PDF devis — conditions",
          "Conditions proposées par défaut sur un nouveau devis (validité, acompte). "
          "Texte contractuel montré au client.")),
        ("pdf:paymentTerms", ("PDF devis / facture", "Titre de la rubrique des modalités de paiement.")),
        ("pdf:structuredCommunication", ("PDF devis / facture", "Communication structurée bancaire (+++…+++). Garder {{value}}.")),
    ]
)
PREFIXE_AUDIT = "pdf:contrat."
CONTEXTE_AUDIT = (
    "PDF du contrat de sous-traitance signé — zone de signature et PISTE D'AUDIT de la signature "
    "électronique (eIDAS). Valeur probante : formulation précise exigée."
)
PRIORITE_2 = [
    ("contrats:sign.", "Page publique de signature (vue par le sous-traitant)",
     "Texte lu par le SOUS-TRAITANT avant et après sa signature (acceptation, valeur de la "
     "signature, langue qui fait foi). Engage l'entreprise : formulation précise exigée."),
    ("contrats:send.subject", "E-mail / WhatsApp au sous-traitant", "Objet du message d'invitation à signer."),
    ("contrats:send.message", "E-mail / WhatsApp au sous-traitant",
     "Message d'invitation à signer envoyé au sous-traitant. Garder {{name}}, {{company}}, {{link}}, {{days}}."),
    ("pdf:mentionReverseCharge", None, None),  # déjà en priorité 1
]

# --- Écrans principaux (namespace / préfixe → écran) -------------------------------------------
ECRANS = [
    ("Connexion", "auth", ""),
    ("Tableau de bord", "dashboard", ""),
    ("Tableau de bord", "sousTraitants", "compliance.alerts."),
    ("Prestations", "sousTraitants", "prestations."),
    ("Prestations", "sousTraitants", "site."),
    ("Conformité", "sousTraitants", "compliance."),
    ("Sous-traitants", "sousTraitants", ""),
    ("Sous-traitants", "statuts", "sousTraitant."),
    ("Sous-traitants", "statuts", "formeJuridique."),
    ("Sous-traitants", "statuts", "metier."),
    ("Sous-traitants", "statuts", "typePrix."),
    ("Sous-traitants", "statuts", "unite."),
    ("Sous-traitants", "statuts", "frequencePaiement."),
    ("Sous-traitants", "statuts", "modePaiement."),
    ("Chantiers", "chantiers", ""),
    ("Contrats", "contrats", ""),
    ("Exports", "exports", ""),
]


def contexte_auto(cle, texte):
    notes = []
    if re.search(r"\{\{\w+\}\}", texte or ""):
        notes.append("Garder les variables {{…}} telles quelles.")
    _, forme = base_pluriel(cle)
    if forme:
        notes.append(f"Forme plurielle « {forme} » (nombre variable).")
    return " ".join(notes)


def lignes():
    fr = {ns: dict(aplatir(charger("fr", ns))) for ns in
          ["pdf", "contrats", "auth", "dashboard", "chantiers", "sousTraitants", "exports", "statuts"]}
    vus = set()
    out = []

    def ajouter(priorite, ecran, ns, cle, contexte, langue, tr):
        id_ = (langue, ns, cle)
        if id_ in vus:
            return
        vus.add(id_)
        texte = tr.get(cle, "")
        src = source_fr(fr[ns], cle)
        ctx = " ".join(x for x in [contexte, contexte_auto(cle, src)] if x)
        out.append([priorite, NOMS_LANGUES[langue], ecran, src, texte, ctx, "À valider", "", "", f"{ns}:{cle}"])

    for langue in LANGUES:
        tr = {ns: dict(aplatir(charger(langue, ns))) for ns in fr}
        pdf_ro_pl = NOTE_PDF_RO_PL if langue in ("ro", "pl") else ""
        # Priorité 1 : mentions légales des PDF + piste d'audit
        for ref, (ecran, ctx) in PRIORITE_1.items():
            ns, cle = ref.split(":")
            ajouter("1 — Mentions légales PDF", ecran, ns, cle, f"{ctx} {pdf_ro_pl}".strip(), langue, tr[ns])
        for cle in tr["pdf"]:
            if f"pdf:{cle}".startswith(PREFIXE_AUDIT):
                ajouter("1 — Mentions légales PDF", "PDF contrat signé — piste d'audit", "pdf", cle,
                        f"{CONTEXTE_AUDIT} {pdf_ro_pl}".strip(), langue, tr["pdf"])
        # Priorité 2 : textes destinés au sous-traitant
        for ref, ecran, ctx in PRIORITE_2:
            if ecran is None:
                continue
            ns, prefixe = ref.split(":")
            for cle in tr[ns]:
                if cle == prefixe or cle.startswith(prefixe):
                    ajouter("2 — Page de signature", ecran, ns, cle, ctx, langue, tr[ns])
        # Priorité 3 : écrans
        for ecran, ns, prefixe in ECRANS:
            for cle in tr[ns]:
                if not prefixe or cle.startswith(prefixe):
                    ajouter("3 — Écrans", ecran, ns, cle, f"Écran « {ecran} ».", langue, tr[ns])
    return out


def construire():
    wb = Workbook()
    ws = wb.active
    ws.title = "Relecture"
    entetes = ["Priorité", "Langue", "Écran ou document", "Texte français", "Traduction actuelle",
               "Contexte", "Statut", "Traduction corrigée", "Commentaire du relecteur", "Clé technique"]
    ws.append(entetes)
    donnees = lignes()
    for l in donnees:
        ws.append(l)

    entete_fill = PatternFill("solid", fgColor="0F172A")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = entete_fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
    largeurs = [24, 16, 34, 50, 50, 55, 20, 50, 40, 38]
    for i, w in enumerate(largeurs, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    couleurs = {"1": "FEE2E2", "2": "FEF3C7"}
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
        p = str(row[0].value)[0]
        if p in couleurs:
            for c in row[:6]:
                c.fill = PatternFill("solid", fgColor=couleurs[p])
        row[9].font = Font(color="94A3B8", size=8)
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:J{len(donnees) + 1}"
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUTS) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"G2:G{len(donnees) + 1}")

    # Synthèse
    syn = wb.create_sheet("Synthèse")
    syn.append(["Langue", "Priorité", "Nombre de textes"])
    compte = {}
    for l in donnees:
        compte[(l[1], l[0])] = compte.get((l[1], l[0]), 0) + 1
    for (langue, prio), n in sorted(compte.items()):
        syn.append([langue, prio, n])
    for c in syn[1]:
        c.font = Font(bold=True)
    syn.column_dimensions["A"].width = 20
    syn.column_dimensions["B"].width = 28
    syn.column_dimensions["C"].width = 18

    # Mode d'emploi
    me = wb.create_sheet("Mode d'emploi", 0)
    texte = [
        "Dossier de relecture native — ConstructFlow",
        f"Généré le {date.today().strftime('%d/%m/%Y')} à partir des fichiers de traduction du logiciel.",
        "",
        "Objectif : faire relire par une personne de langue maternelle les textes néerlandais, roumains et polonais.",
        "",
        "Comment remplir l'onglet « Relecture » :",
        "1. Filtrez par « Langue » (une feuille par relecteur si besoin).",
        "2. Commencez par la priorité 1 (mentions légales des PDF, en rouge), puis 2 (page de signature, en orange), puis 3 (écrans).",
        "3. Pour chaque ligne : lisez le texte français et la traduction actuelle, puis choisissez le statut :",
        "   - « Validé » si la traduction est correcte ;",
        "   - « Correction proposée » et écrivez la bonne traduction dans « Traduction corrigée ».",
        "4. Les éléments entre doubles accolades, par exemple {{count}} ou {{date}}, sont remplacés automatiquement : gardez-les tels quels.",
        "5. Les formes plurielles (one, few, many, other) sont des variantes selon le nombre : traduisez chacune.",
        "6. Ne modifiez pas la colonne « Clé technique » : elle permet de réintégrer vos corrections automatiquement.",
        "",
        "Attention : les mentions légales (autoliquidation, TVA 6 %, piste d'audit) engagent l'entreprise. "
        "En cas de doute sur une formulation légale, signalez-le dans le commentaire plutôt que de la reformuler.",
        "Les contrats et devis ne sont produits qu'en français, néerlandais ou anglais : pour le roumain et le polonais, "
        "les mentions de priorité 1 n'apparaissent pas aujourd'hui sur un document client.",
    ]
    for t in texte:
        me.append([t])
    me["A1"].font = Font(bold=True, size=14)
    me.column_dimensions["A"].width = 140
    for row in me.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    os.makedirs(os.path.dirname(SORTIE), exist_ok=True)
    wb.save(SORTIE)
    return donnees, compte


if __name__ == "__main__":
    donnees, compte = construire()
    print(f"{len(donnees)} lignes -> {os.path.relpath(SORTIE, RACINE)}")
    for (langue, prio), n in sorted(compte.items()):
        print(f"  {langue:18} {prio:28} {n}")
